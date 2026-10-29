from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from apps.clusters.models import Cluster
from apps.geography.models import District, Province


PROVINCE_ALIASES = {
    "MANICALAND": "Manicaland",
    "MASHCENTRAL": "Mashonaland Central",
    "MASHEAST": "Mashonaland East",
    "MASHWEST": "Mashonaland West",
    "MATABELELAND SOUTH": "Matabeleland South",
    "MIDLANDS": "Midlands",
    "MIDLADS": "Midlands",
}


class Command(BaseCommand):
    help = "Import cluster names and district assignments from the SGR workbook."

    def add_arguments(self, parser):
        parser.add_argument(
            "--workbook",
            type=Path,
            default=Path(__file__).resolve().parents[5] / "SGR SOURCE DATABASE.xlsx",
        )

    def handle(self, *args, **options):
        workbook_path = options["workbook"]
        if not workbook_path.is_file():
            raise CommandError(f"Workbook not found: {workbook_path}")

        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        if "CLUSTER PERFORMANCE" not in workbook.sheetnames:
            raise CommandError("Workbook is missing the CLUSTER PERFORMANCE sheet.")

        rows = workbook["CLUSTER PERFORMANCE"].iter_rows(values_only=True)
        header = None
        imported = 0
        skipped = 0
        with transaction.atomic():
            for row in rows:
                values = list(row)
                if header is None:
                    normalized = [str(value or "").strip().upper() for value in values]
                    if "CLUSTER NAME" in normalized:
                        header = {name: index for index, name in enumerate(normalized)}
                    continue

                name = str(values[header["CLUSTER NAME"]] or "").strip()
                province_label = str(values[header["PROVINCE"]] or "").strip().upper()
                district_label = str(values[header["DISTRICT"]] or "").strip()
                province_name = PROVINCE_ALIASES.get(province_label)
                if not name or not province_name or not district_label or district_label.upper() in {"NK", "UNKNOWN"}:
                    skipped += 1
                    continue

                province = Province.objects.filter(name__iexact=province_name).first()
                district = District.objects.filter(
                    province=province,
                    name__iexact=district_label,
                    active=True,
                ).first() if province else None
                if district is None:
                    skipped += 1
                    continue

                cluster = Cluster.objects.filter(name__iexact=name).first()
                if cluster is None:
                    Cluster.objects.create(name=name, province=province, district=district)
                else:
                    changed_fields = []
                    if cluster.province_id != province.pk:
                        cluster.province = province
                        changed_fields.append("province")
                    if cluster.district_id != district.pk:
                        cluster.district = district
                        changed_fields.append("district")
                    if not cluster.active:
                        cluster.active = True
                        changed_fields.append("active")
                    if changed_fields:
                        cluster.save(update_fields=[*changed_fields, "updated_at"])
                imported += 1

        workbook.close()
        self.stdout.write(self.style.SUCCESS(f"Loaded {imported} cluster records; skipped {skipped} rows with unresolved locations."))