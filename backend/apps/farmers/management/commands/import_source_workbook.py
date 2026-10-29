from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from apps.farmers.models import Farm, Farmer
from apps.finance.models import Asset, Expense, PayrollEntry, Receipt
from apps.production.models import ProductionPlan
from apps.recovery.models import RecoveryRecord
from apps.subscriptions.models import FarmerSubscription, Subscription
from apps.geography.models import District, Province, Ward
from apps.source_data.models import (
    BalanceSheetLine,
    BudgetLine,
    ClusterPerformance,
    DataBankEntry,
    FinanceCashFlow,
    MonthlySalary,
    ProfitLossLine,
)


SEASON = "2026/2027"


def text(value):
    if value is None:
        return ""
    return " ".join(str(value).replace("\xa0", " ").split()).strip()


def decimal(value):
    if value in (None, "", "-", "-   ", " -   ", "\xa0"):
        return Decimal("0")
    try:
        return Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, ValueError):
        return Decimal("0")


def date_value(value):
    return value.date() if hasattr(value, "date") else value


def name_parts(full_name):
    parts = text(full_name).split()
    if not parts:
        return "Legacy", "", "Farmer"
    if len(parts) == 1:
        return parts[0], "", parts[0]
    return parts[0], " ".join(parts[1:-1]), parts[-1]


def status_for_recovery(status):
    value = text(status).upper()
    if "CLEARED" in value:
        return RecoveryRecord.STATUS_PAID
    if "PARTIAL" in value:
        return RecoveryRecord.STATUS_IN_PROGRESS
    if "NOT DELIVERED" in value:
        return RecoveryRecord.STATUS_PENDING
    return RecoveryRecord.STATUS_PENDING


class Command(BaseCommand):
    help = "Import the SGR source workbook into the configured Django database."

    def add_arguments(self, parser):
        parser.add_argument("workbook", type=Path)

    @transaction.atomic
    def handle(self, *args, **options):
        workbook_path = options["workbook"]
        if not workbook_path.exists():
            raise CommandError(f"Workbook not found: {workbook_path}")

        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        counts = {
            "farmers": 0,
            "farms": 0,
            "subscriptions": 0,
            "subscription_plans": 0,
            "recovery_records": 0,
            "receipts": 0,
            "expenses": 0,
            "assets": 0,
            "payroll_entries": 0,
            "production_plans": 0,
            "data_bank": 0,
            "cluster_performance": 0,
            "finance_cash_flow": 0,
            "balance_sheet": 0,
            "profit_loss": 0,
            "monthly_salary": 0,
            "budget_lines": 0,
        }
        farmers_by_id = {}

        pipeline = workbook["DATA_PIPELINE"]
        for row_number, row in enumerate(pipeline.iter_rows(min_row=6, values_only=True), start=6):
            if not text(row[1]) and not text(row[4]):
                continue
            farmer_name = text(row[1]) or "Legacy Farmer"
            national_id = text(row[4]) or f"LEGACY-{row_number:06d}"
            first_name, middle_name, surname = name_parts(farmer_name)
            province_name = text(row[7]) or "Unknown"
            district_name = text(row[6]) or "Unknown"
            province, _ = Province.objects.get_or_create(
                name=province_name,
                defaults={"code": province_name[:10].upper().replace(" ", "-")},
            )
            district, _ = District.objects.get_or_create(province=province, name=district_name)
            farmer, _ = Farmer.objects.update_or_create(
                national_id=national_id,
                defaults={
                    "first_name": first_name,
                    "middle_name": middle_name,
                    "surname": surname,
                    "phone_primary": text(row[5]),
                    "province": province,
                    "district": district,
                    "verification_status": Farmer.STATUS_VERIFIED if text(row[9]).upper() == "EXISTING" else Farmer.STATUS_SUBMITTED,
                },
            )
            farmers_by_id[national_id] = farmer
            counts["farmers"] += 1

            farm_name = text(row[2]) or f"{farmer.full_name} Farm"
            Farm.objects.update_or_create(
                farmer=farmer,
                farm_name=farm_name,
                defaults={"total_hectares": decimal(row[19])},
            )
            counts["farms"] += 1

            category = text(row[20]) or "Uncategorised"
            plan, created = Subscription.objects.get_or_create(
                code=f"LEGACY-{category.upper().replace(' ', '-')}",
                defaults={"name": category, "fee": decimal(row[21])},
            )
            if created:
                counts["subscription_plans"] += 1
            FarmerSubscription.objects.update_or_create(
                farmer=farmer,
                season=SEASON,
                defaults={
                    "category": category,
                    "expected_amount": decimal(row[21]),
                    "paid_amount": decimal(row[22]),
                    "outstanding_amount": decimal(row[23]),
                    "payment_status": text(row[24]),
                },
            )
            counts["subscriptions"] += 1

        debt = workbook["INPUT DEBT"]
        for row_number, row in enumerate(debt.iter_rows(min_row=6, values_only=True), start=6):
            national_id = text(row[4])
            lookup_id = national_id or f"LEGACY-DEBT-{text(row[0]) or row_number}"
            farmer = farmers_by_id.get(lookup_id)
            if farmer is None:
                farmer = Farmer.objects.filter(national_id=lookup_id).first()
            if farmer is None:
                first_name, middle_name, surname = name_parts(text(row[3]))
                province_name = text(row[1]) or "Unknown"
                province, _ = Province.objects.get_or_create(
                    name=province_name,
                    defaults={"code": province_name[:10].upper().replace(" ", "-")},
                )
                farmer = Farmer.objects.create(
                    national_id=lookup_id,
                    first_name=first_name,
                    middle_name=middle_name,
                    surname=surname,
                    province=province,
                )
                farmers_by_id[lookup_id] = farmer
                counts["farmers"] += 1
            defaults = {
                "province": text(row[1]),
                "farm_name": text(row[2]),
                "crop": text(row[5]),
                "contracted_hectares": decimal(row[6]),
                "arda_grand_total": decimal(row[7]),
                "debt_amount": decimal(row[9]),
                "grain_recovery_tonnes": decimal(row[12]),
                "grain_delivered_tonnes": decimal(row[13]),
                "outstanding_tonnes": decimal(row[14]),
                "status": status_for_recovery(row[15]),
            }
            record = RecoveryRecord.objects.filter(
                farmer=farmer,
                farm_name=defaults["farm_name"],
                crop=defaults["crop"],
                contracted_hectares=defaults["contracted_hectares"],
            ).first()
            if record:
                for key, value in defaults.items():
                    setattr(record, key, value)
                record.save()
            else:
                RecoveryRecord.objects.create(farmer=farmer, **defaults)
            counts["recovery_records"] += 1

        receipts = workbook["DATA_RECEIPTS"]
        for row_number, row in enumerate(receipts.iter_rows(min_row=6, values_only=True), start=6):
            if not text(row[2]):
                continue
            Receipt.objects.update_or_create(
                source_row=row_number,
                defaults={
                    "month": date_value(row[0]),
                    "receipt_date": date_value(row[1]),
                    "farmer_name": text(row[2]),
                    "amount_due": decimal(row[3]),
                    "amount_paid": decimal(row[4]),
                    "balance": decimal(row[5]),
                    "payment_type": text(row[6]),
                },
            )
            counts["receipts"] += 1

        expenses = workbook["DATA_EXPENSES"]
        for row_number, row in enumerate(expenses.iter_rows(min_row=6, values_only=True), start=6):
            if not text(row[1]):
                continue
            Expense.objects.update_or_create(
                source_row=row_number,
                defaults={
                    "expense_date": date_value(row[0]),
                    "particulars": text(row[1]),
                    "amount_debit": decimal(row[2]),
                    "reimbursement": decimal(row[3]),
                    "net_expense": decimal(row[4]),
                },
            )
            counts["expenses"] += 1

        assets = workbook["SGR ASSET REGISTER"]
        for row in assets.iter_rows(min_row=7, values_only=True):
            if not text(row[0]) or not text(row[1]):
                continue
            Asset.objects.update_or_create(
                asset_id=text(row[0]),
                defaults={
                    "description": text(row[1]),
                    "category": text(row[2]),
                    "purchase_date": date_value(row[3]),
                    "cost": decimal(row[4]),
                    "residual_value": decimal(row[5]),
                    "useful_life_years": decimal(row[6]),
                    "annual_depreciation": decimal(row[7]),
                    "accumulated_depreciation": decimal(row[8]),
                    "closing_nbv": decimal(row[9]),
                },
            )
            counts["assets"] += 1

        payroll = workbook["SALARY SCHEDULE"]
        for row in payroll.iter_rows(min_row=7, values_only=True):
            if not text(row[0]) or not text(row[1]):
                continue
            PayrollEntry.objects.update_or_create(
                employee_id=text(row[0]),
                defaults={
                    "employee_position": text(row[1]),
                    "department": text(row[2]),
                    "basic_salary_monthly": decimal(row[3]),
                    "housing_allowance_monthly": decimal(row[4]),
                    "transport_allowance_monthly": decimal(row[5]),
                    "other_allowance_monthly": decimal(row[6]),
                    "total_monthly_cost": decimal(row[7]),
                    "annual_salary_cost": decimal(row[8]),
                    "annual_employment_cost": decimal(row[9]),
                },
            )
            counts["payroll_entries"] += 1

        production = workbook["PRODUCTION"]
        for row in production.iter_rows(min_row=11, max_row=12, values_only=True):
            if not text(row[0]):
                continue
            ProductionPlan.objects.update_or_create(
                crop=text(row[0]),
                season=SEASON,
                defaults={
                    "target_hectares": decimal(row[1]),
                    "actual_hectares": decimal(row[2]),
                    "variance_hectares": decimal(row[3]),
                    "achievement": decimal(row[4]),
                    "remaining_hectares": decimal(row[5]),
                },
            )
            counts["production_plans"] += 1

        for row_number, row in enumerate(workbook["DATA_BANK"].iter_rows(min_row=6, max_row=24, values_only=True), start=6):
            if not text(row[0]):
                continue
            DataBankEntry.objects.update_or_create(
                source_row=row_number,
                defaults={
                    "month": text(row[0]),
                    "opening_balance": decimal(row[1]),
                    "receipts": decimal(row[2]),
                    "expenses": decimal(row[3]),
                    "closing_balance": decimal(row[4]),
                    "source_check": text(row[5]),
                    "difference": decimal(row[6]),
                    "status": text(row[7]),
                },
            )
            counts["data_bank"] += 1

        for row in workbook["CLUSTER PERFORMANCE"].iter_rows(min_row=9, values_only=True):
            if not text(row[0]):
                continue
            ClusterPerformance.objects.update_or_create(
                cluster_name=text(row[0]),
                defaults={
                    "farmers_managed": int(decimal(row[1])),
                    "maize_hectares": decimal(row[2]),
                    "sorghum_hectares": decimal(row[3]),
                    "total_hectares": decimal(row[4]),
                    "paid_farmers": int(decimal(row[5])),
                    "unpaid_farmers": int(decimal(row[6])),
                    "payment_percent": decimal(row[7]),
                    "province": text(row[8]),
                    "district": text(row[9]),
                    "cluster_contact": text(row[10]),
                    "contact_details": text(row[11]),
                    "expected_subscriptions": decimal(row[12]),
                    "amount_paid": decimal(row[13]),
                },
            )
            counts["cluster_performance"] += 1

        for row in workbook["FINANCE"].iter_rows(min_row=11, max_row=20, values_only=True):
            if not text(row[0]):
                continue
            FinanceCashFlow.objects.update_or_create(
                month=text(row[0]),
                defaults={
                    "opening_balance": decimal(row[1]),
                    "receipts": decimal(row[2]),
                    "expenses": decimal(row[3]),
                    "closing_balance": decimal(row[4]),
                },
            )
            counts["finance_cash_flow"] += 1

        balance_sections = {6: "ASSETS", 12: "LIABILITIES", 16: "EQUITY"}
        current_section = ""
        for row_number, row in enumerate(workbook["BALANCE SHEET"].iter_rows(min_row=6, max_row=22, values_only=True), start=6):
            if row_number in balance_sections:
                current_section = balance_sections[row_number]
                continue
            if not text(row[0]) or row[1] is None:
                continue
            BalanceSheetLine.objects.update_or_create(
                source_row=row_number,
                defaults={"section": current_section, "particulars": text(row[0]), "amount": decimal(row[1])},
            )
            counts["balance_sheet"] += 1

        for row_number, row in enumerate(workbook["PROFIT & LOSS"].iter_rows(min_row=8, max_row=21, values_only=True), start=8):
            if not text(row[0]) or row_number in (11,):
                continue
            ProfitLossLine.objects.update_or_create(
                source_row=row_number,
                defaults={
                    "particulars": text(row[0]),
                    "june": decimal(row[1]),
                    "july": decimal(row[2]),
                    "august": decimal(row[3]),
                    "september": decimal(row[4]),
                    "october": decimal(row[5]),
                    "november": decimal(row[6]),
                    "december": decimal(row[7]),
                    "jun_aug_total": decimal(row[8]),
                },
            )
            counts["profit_loss"] += 1

        for row in workbook["MONTHLY SALARY"].iter_rows(min_row=5, max_row=11, values_only=True):
            if not text(row[0]):
                continue
            MonthlySalary.objects.update_or_create(
                month=text(row[0]),
                defaults={
                    "planned_salary": decimal(row[1]),
                    "actual_salary": decimal(row[2]),
                    "variance": decimal(row[3]),
                    "status": text(row[4]),
                },
            )
            counts["monthly_salary"] += 1

        for budget_name, sheet_name, start_row, end_row in [
            ("MAIZE", "MAIZE BUDGET", 9, 48),
            ("SORGHUM", "SORGHUM BUDGET", 10, 66),
        ]:
            for row_number, row in enumerate(workbook[sheet_name].iter_rows(min_row=start_row, max_row=end_row, values_only=True), start=start_row):
                item = text(row[1] if budget_name == "SORGHUM" else row[0])
                if not item or item.upper() in {"TOTAL", "SUBTOTAL", "TOTAL VARIABLE COST"}:
                    continue
                group = text(row[0] if budget_name == "SORGHUM" else "")
                BudgetLine.objects.update_or_create(
                    budget=budget_name,
                    source_row=row_number,
                    defaults={
                        "group": group,
                        "item": item,
                        "quantity": decimal(row[1] if budget_name == "MAIZE" else row[3]),
                        "rate": decimal(row[2] if budget_name == "MAIZE" else row[6]),
                        "unit": text(row[4] if budget_name == "MAIZE" else row[4]),
                        "cost_per_hectare": decimal(row[3] if budget_name == "MAIZE" else row[7]),
                        "total_cost": decimal(row[4] if budget_name == "MAIZE" else row[8]),
                        "funded_cost": decimal(row[6] if budget_name == "MAIZE" else 0),
                    },
                )
                counts["budget_lines"] += 1

        self.stdout.write(self.style.SUCCESS("Workbook import completed."))
        for key, value in counts.items():
            self.stdout.write(f"{key}: {value}")
